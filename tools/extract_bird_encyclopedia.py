#!/usr/bin/env python3
"""
한반도 조류 598종 고화질 사진도감(xlsx)에서
1) 598종 고화질 원본 사진 -> app/src/main/assets/birds/{num}.jpg
2) 도감 목록용 최적화 썸네일 -> app/src/main/assets/birds_thumb/{num}.jpg
3) 598종 상세 설명 및 메타데이터 -> app/src/main/assets/birds_encyclopedia.json
을 추출하여 생성합니다.
"""
import io
import json
import os
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from PIL import Image
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "한반도_조류_598종_고화질_사진도감_2025.xlsx"
ASSETS_DIR = ROOT / "app/src/main/assets"
BIRDS_IMG_DIR = ASSETS_DIR / "birds"
BIRDS_THUMB_DIR = ASSETS_DIR / "birds_thumb"
JSON_OUT_PATH = ASSETS_DIR / "birds_encyclopedia.json"


def main():
    if not XLSX_PATH.exists():
        print(f"Error: {XLSX_PATH} not found!")
        sys.exit(1)

    BIRDS_IMG_DIR.mkdir(parents=True, exist_ok=True)
    BIRDS_THUMB_DIR.mkdir(parents=True, exist_ok=True)

    print("1. 엑셀 워크북 로딩 중...")
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["01_전체종_사진"]

    print("2. 엑셀 드로잉 및 이미지 매핑 파싱 중...")
    with zipfile.ZipFile(XLSX_PATH, "r") as z:
        # drawing1.xml.rels
        rels_xml = z.read("xl/drawings/_rels/drawing1.xml.rels")
        rels_tree = ET.fromstring(rels_xml)
        rel_map = {
            r.attrib["Id"]: r.attrib["Target"]
            for r in rels_tree.findall("{http://schemas.openxmlformats.org/package/2006/relationships}Relationship")
        }

        # drawing1.xml
        draw_xml = z.read("xl/drawings/drawing1.xml")
        draw_tree = ET.fromstring(draw_xml)
        row_to_img = {}
        for anc in draw_tree:
            from_el = anc.find("{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}from")
            if from_el is not None:
                r = int(from_el.find("{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}row").text)
                blip = anc.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
                if blip is not None:
                    embed = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    target = rel_map.get(embed)
                    if target:
                        # target can be /xl/media/image1.jpeg or ../media/image1.jpeg
                        target_clean = target.lstrip("/").replace("../", "xl/")
                        if not target_clean.startswith("xl/"):
                            target_clean = "xl/" + target_clean
                        row_to_img[r] = target_clean

        print(f"매핑된 사진 수: {len(row_to_img)}")

        birds = []
        print("3. 598종 사진 추출, 썸네일 생성 및 메타데이터 정리 중...")
        for idx in range(1, 599):
            excel_row = idx + 4
            zero_row = excel_row - 1

            num = int(ws.cell(excel_row, 1).value)
            order = str(ws.cell(excel_row, 2).value or "").strip()
            family = str(ws.cell(excel_row, 3).value or "").strip()
            name = str(ws.cell(excel_row, 4).value or "").strip()
            sci = str(ws.cell(excel_row, 5).value or "").strip()
            eng = str(ws.cell(excel_row, 6).value or "").strip()
            cat = str(ws.cell(excel_row, 8).value or "").strip()
            desc = str(ws.cell(excel_row, 9).value or "").strip()
            desc_src = str(ws.cell(excel_row, 10).value or "").strip()
            eng_desc = str(ws.cell(excel_row, 11).value or "").strip()
            subs_raw = ws.cell(excel_row, 13).value or ""
            subs = [s.strip() for s in str(subs_raw).split("\n") if s.strip() and s.strip() != "—"]
            status = str(ws.cell(excel_row, 16).value or "").strip()
            if status == "미제공":
                status = ""
            obs_val = ws.cell(excel_row, 17).value
            obs = obs_val if isinstance(obs_val, int) else 0
            license_str = str(ws.cell(excel_row, 19).value or "").strip()
            author = str(ws.cell(excel_row, 21).value or "").strip()

            # 이미지 추출
            img_zip_path = row_to_img.get(zero_row)
            if not img_zip_path or img_zip_path not in z.namelist():
                # Fallback: image{idx}.jpeg
                img_zip_path = f"xl/media/image{idx}.jpeg"

            img_bytes = z.read(img_zip_path)
            # 원본 저장
            out_img_path = BIRDS_IMG_DIR / f"{num}.jpg"
            out_img_path.write_bytes(img_bytes)

            # 썸네일 생성 (가로 160px 비례 축소)
            try:
                pil_img = Image.open(io.BytesIO(img_bytes))
                # RGB로 변환 (RGBA/CMYK 등 대응)
                if pil_img.mode != "RGB":
                    pil_img = pil_img.convert("RGB")
                w, h = pil_img.size
                thumb_w = 160
                thumb_h = max(1, int(h * (thumb_w / w)))
                thumb_img = pil_img.resize((thumb_w, thumb_h), Image.Resampling.LANCZOS)
                thumb_img.save(BIRDS_THUMB_DIR / f"{num}.jpg", format="JPEG", quality=85)
            except Exception as e:
                print(f"썸네일 생성 실패 (num={num}): {e}")

            birds.append(
                {
                    "num": num,
                    "name": name,
                    "sci": sci,
                    "eng": eng,
                    "order": order,
                    "family": family,
                    "cat": cat,
                    "desc": desc,
                    "descSrc": desc_src,
                    "engDesc": eng_desc,
                    "subs": subs,
                    "status": status,
                    "obs": obs,
                    "license": license_str,
                    "author": author,
                }
            )

            if idx % 100 == 0 or idx == 598:
                print(f"  [{idx}/598] {name} 완료")

    print("4. birds_encyclopedia.json 저장 중...")
    with open(JSON_OUT_PATH, "w", encoding="utf-8") as fp:
        json.dump(birds, fp, ensure_ascii=False, indent=2)

    print(f"완료! 원본 {len(birds)}장 및 썸네일 {len(birds)}장, 메타데이터 JSON 생성 완료.")


if __name__ == "__main__":
    main()
